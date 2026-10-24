from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ai_asset_management_company_blueprint.md"
OUTPUT = ROOT / "ai_asset_management_company_blueprint.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([run_properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text):
    token_re = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\(https?://[^)]+\))")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            link_match = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = RGBColor(35, 35, 35)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 18, 7),
        ("Heading 2", 13, 13, 5),
        ("Heading 3", 11, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Menlo"
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor(45, 45, 45)
    code_style.paragraph_format.left_indent = Inches(0.28)
    code_style.paragraph_format.space_before = Pt(5)
    code_style.paragraph_format.space_after = Pt(8)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI Asset Management Company Blueprint")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(110, 110, 110)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)
    add_footer(section)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    first_heading = True
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                p.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("| ") and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
            rows = []
            header = [cell.strip() for cell in line.strip("|").split("|")]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([cell.strip() for cell in lines[i].strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.autofit = False
            table.columns[0].width = Inches(1.65)
            if len(header) > 1:
                table.columns[1].width = Inches(4.95)
            for j, value in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = value
                set_cell_shading(cell, "1F4E79")
                set_cell_borders(cell)
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
            set_repeat_table_header(table.rows[0])
            for row_index, values in enumerate(rows):
                cells = table.add_row().cells
                for j, value in enumerate(values):
                    cells[j].text = value
                    set_cell_borders(cells[j])
                    set_cell_margins(cells[j])
                    cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if row_index % 2:
                        set_cell_shading(cells[j], "F3F6F9")
                    for run in cells[j].paragraphs[0].runs:
                        run.font.size = Pt(9.3)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title" if first_heading else "Heading 1")
            add_rich_text(p, line[2:])
            first_heading = False
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_rich_text(p, line[3:])
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_rich_text(p, line[4:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            add_rich_text(p, line)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.32)
            p.paragraph_format.right_indent = Inches(0.22)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:])
            run.bold = True
            run.italic = True
            run.font.size = Pt(11.2)
        elif line.startswith("*") and line.endswith("*"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(90, 90, 90)
        else:
            p = doc.add_paragraph()
            add_rich_text(p, line)
        i += 1

    props = doc.core_properties
    props.title = "AI Asset Management Company Blueprint"
    props.subject = "Strategy, product, governance, business model, and launch roadmap"
    props.keywords = "AI, asset management, agents, investment research, portfolio intelligence"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
